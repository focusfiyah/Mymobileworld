#!/usr/bin/env python3
"""Grace's sales tracker: AI videos that made a sale (Ralph/Grace, 2026-10-20).

Grace sends a video link + how many sales it got. This pulls everything TikTok shows publicly; Claude fills the
judgment fields (pain point, problem it solves, why it converted, faceless, format, on-screen text) after looking
at the frames, then exports the sheet.

  python3 grace/winners/winners.py pull <link> --sales N [--note "..."]   # add or update one video
  python3 grace/winners/winners.py set <id> field=value [field=value ...]  # fill a judgment field
  python3 grace/winners/winners.py refresh                                  # re-pull views/saves for every row
  python3 grace/winners/winners.py export                                   # winners.csv + winners.xlsx
  python3 grace/winners/winners.py todo                                     # rows with empty judgment fields

Source of truth: winners.json (this folder). Media + frames go to the scratchpad (--out), never the repo.
"""
import argparse, csv, difflib, json, os, re, subprocess, sys, wave
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(os.path.dirname(HERE))
DB = os.path.join(HERE, "winners.json")
TT = os.path.join(REPO, ".claude/skills/tiktok-shop-coach/scripts/tiktok.py")
HOOKS = os.path.join(REPO, "grace/hooks/hooks.json")

AUTO = ["id", "url", "share_link", "account", "followers", "posted", "duration_s", "views", "likes", "comments",
        "shares", "saves", "product", "product_category", "product_url", "caption", "script", "script_source",
        "hook", "hook_id", "hook_type", "hook_match", "pulled"]
JUDGED = ["faceless", "format", "onscreen_text", "pain_point", "problem_solved", "why_converted", "job"]
# Sheet column order (what Grace reads first goes left).
COLUMNS = [("posted", "Posted"), ("account", "Account"), ("product", "Product"), ("sales", "Sales"),
           ("views", "Views"), ("views_per_sale", "Views per sale"), ("faceless", "Faceless"),
           ("format", "Format"), ("hook", "Hook (first line said)"), ("hook_type", "Hook type"),
           ("hook_id", "Library hook"), ("onscreen_text", "On-screen text"), ("pain_point", "Pain point"),
           ("problem_solved", "Problem the product solves"), ("why_converted", "Why it converted (our read)"),
           ("script", "Script used (as posted)"), ("url", "Link to video"), ("product_url", "Product link"),
           ("saves", "Saves"), ("shares", "Shares"), ("duration_s", "Length (s)"), ("job", "Our job folder"),
           ("note", "Note")]


def load():
    if not os.path.exists(DB):
        return {"updated": None, "videos": []}
    return json.load(open(DB))


def save(db):
    db["updated"] = date.today().isoformat()
    for v in db["videos"]:
        s = v.get("sales") or 0
        v["views_per_sale"] = round(v["views"] / s) if s and v.get("views") is not None else None
    db["videos"].sort(key=lambda v: v.get("posted") or "", reverse=True)
    json.dump(db, open(DB, "w"), indent=1, ensure_ascii=False)


def first_sentence(text):
    m = re.split(r"(?<=[.!?])\s+", text.strip(), maxsplit=1)
    return m[0] if m else text


def match_hook(line):
    """Closest hook in Grace's library. Blanks (__) match anything, so compare the fixed words only."""
    hooks = json.load(open(HOOKS))["hooks"]
    low = line.lower()
    best, score = None, 0.0
    for h in hooks:
        parts = [p.strip().lower() for p in h["hook"].split("__") if p.strip()]
        fixed = " ".join(parts)
        r = difflib.SequenceMatcher(None, low[:len(fixed) + 40], fixed).ratio()
        if parts and all(p in low for p in parts):
            r = max(r, 0.95)
        if r > score:
            best, score = h, r
    return best, round(score, 2)


def local_stt(mp4, scratch):
    """Free transcript when TikTok has no captions: faster-whisper small.en on CPU (pip install faster-whisper)."""
    wav = os.path.join(scratch, "audio16k.wav")
    subprocess.run(["ffmpeg", "-hide_banner", "-loglevel", "error", "-y", "-i", mp4, "-ac", "1", "-ar", "16000", wav],
                   check=True)
    import numpy as np
    from faster_whisper import WhisperModel
    w = wave.open(wav)
    a = np.frombuffer(w.readframes(w.getnframes()), np.int16).astype(np.float32) / 32768
    segs, _ = WhisperModel("small.en", device="cpu", compute_type="int8").transcribe(a, beam_size=5)
    return " ".join(s.text.strip() for s in segs)


def pull(link, out, sales=None, note=None):
    r = subprocess.run([sys.executable, TT, "video", link, "--out", out], capture_output=True, text=True)
    m = re.search(r"-> (\S+)", r.stdout)
    if r.returncode or not m:
        sys.exit(f"tiktok.py failed for {link}:\n{r.stdout[-800:]}\n{r.stderr[-800:]}")
    folder = m.group(1)
    meta = json.load(open(os.path.join(folder, "meta.json")))
    tpath = os.path.join(folder, "transcript.txt")
    if os.path.exists(tpath):
        script, source = open(tpath).read().strip(), "TikTok captions"
    else:
        script, source = local_stt(os.path.join(folder, "video.mp4"), folder), "local speech-to-text (faster-whisper)"
    hook = first_sentence(script) if script else ""
    h, score = match_hook(hook) if hook else (None, 0)
    p = meta.get("product") or {}
    row = {
        "id": meta["id"], "url": meta["url"], "share_link": link, "account": "@" + meta["author"],
        "followers": meta.get("author_followers"), "posted": meta.get("posted"), "duration_s": meta.get("duration_s"),
        "views": meta.get("views"), "likes": meta.get("likes"), "comments": meta.get("comments"),
        "shares": meta.get("shares"), "saves": meta.get("saves"), "product": p.get("title"),
        "product_category": p.get("category"), "product_url": (p.get("url") or "").split("?")[0],
        "caption": meta.get("caption"), "script": script, "script_source": source, "hook": hook,
        "hook_id": h["id"] if h and score >= 0.8 else "own line",
        "hook_type": h["type"] if h and score >= 0.8 else "",
        "hook_match": f"closest {h['id']} ({h['type']}) '{h['hook']}' score {score}" if h else "",
        "pulled": date.today().isoformat(), "frames": os.path.join(folder, "sheet.png"),
    }
    db = load()
    old = next((v for v in db["videos"] if v["id"] == row["id"]), None)
    if old:
        keep = {k: old[k] for k in old if k not in AUTO}   # never overwrite Grace's sales or our judgment
        old.update(row); old.update(keep)
        if old.get("hook_id_override"):
            old["hook_id"], old["hook_type"] = old["hook_id_override"], old.get("hook_type_override", old["hook_type"])
        v = old
    else:
        v = dict(row, **{k: "" for k in JUDGED})
        db["videos"].append(v)
    if sales is not None:
        v["sales"] = sales
    if note:
        v["note"] = note
    save(db)
    print(f"OK {v['id']} {v['product'][:60]!r} views={v['views']} sales={v.get('sales')} hook={v['hook_id']}")
    print(f"   frames: {row['frames']}  (look at them, then fill: {', '.join(k for k in JUDGED if not v.get(k))})")
    return v


def export():
    db = load()
    rows = [[("" if v.get(k) is None else v.get(k)) for k, _ in COLUMNS] for v in db["videos"]]
    with open(os.path.join(HERE, "winners.csv"), "w", newline="") as f:
        w = csv.writer(f); w.writerow([c for _, c in COLUMNS]); w.writerows(rows)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook(); ws = wb.active; ws.title = "Sales videos"
    ws.append([c for _, c in COLUMNS])
    for r in rows:
        ws.append(r)
    widths = {"Script used (as posted)": 70, "Why it converted (our read)": 55, "Pain point": 35,
              "Problem the product solves": 35, "Hook (first line said)": 40, "Product": 40, "Link to video": 30}
    for i, (_, c) in enumerate(COLUMNS, 1):
        col = ws.cell(1, i).column_letter
        ws.column_dimensions[col].width = widths.get(c, 14)
        ws.cell(1, i).font = Font(bold=True, color="FFFFFF")
        ws.cell(1, i).fill = PatternFill("solid", fgColor="E8590C")
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "D2"
    wb.save(os.path.join(HERE, "winners.xlsx"))
    print(f"OK export {len(rows)} rows -> winners.csv, winners.xlsx")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("pull"); p.add_argument("links", nargs="+"); p.add_argument("--sales", type=int)
    p.add_argument("--note"); p.add_argument("--out", default=os.environ.get("WINNERS_OUT", "/tmp/winners"))
    s = sub.add_parser("set"); s.add_argument("id"); s.add_argument("pairs", nargs="+")
    r = sub.add_parser("refresh"); r.add_argument("--out", default=os.environ.get("WINNERS_OUT", "/tmp/winners"))
    sub.add_parser("export"); sub.add_parser("todo")
    a = ap.parse_args()
    if a.cmd == "pull":
        for l in a.links:
            pull(l, a.out, a.sales, a.note)
    elif a.cmd == "refresh":
        for v in load()["videos"]:
            pull(v["share_link"], a.out)
    elif a.cmd == "set":
        db = load(); v = next(x for x in db["videos"] if x["id"] == a.id)
        for kv in a.pairs:
            k, val = kv.split("=", 1)
            v[k] = int(val) if k == "sales" else val
        save(db); print("OK set", a.id, ", ".join(kv.split("=")[0] for kv in a.pairs))
    elif a.cmd == "export":
        export()
    elif a.cmd == "todo":
        for v in load()["videos"]:
            miss = [k for k in JUDGED + ["sales"] if v.get(k) in ("", None)]
            if miss:
                print(v["id"], v["product"][:40], "->", ", ".join(miss))


if __name__ == "__main__":
    main()
